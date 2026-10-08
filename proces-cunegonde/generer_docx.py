# -*- coding: utf-8 -*-
"""Word script of the trial: every line numbered, timed, with its entry cue."""
import re, sys
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
WPM = 140.0                      # pace assumed for a spoken courtroom register
TRANSITION = 1.5                 # seconds lost between two speakers

# ---------------- minimal DOM ----------------
class Node:
    __slots__ = ("tag","attrs","children","text","parent")
    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.children, self.text, self.parent = tag, attrs or {}, [], None, parent
    def cls(self): return (self.attrs.get("class") or "").split()
    def find_all(self, pred, out=None):
        out = [] if out is None else out
        for c in self.children:
            if c.tag and pred(c): out.append(c)
            if c.tag: c.find_all(pred, out)
        return out
    def first(self, pred):
        r = self.find_all(pred); return r[0] if r else None

VOID = {"br","link","meta","img","hr","input"}
class Build(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True); self.root = Node(None); self.cur = self.root
    def handle_starttag(self, tag, attrs):
        n = Node(tag, dict(attrs), self.cur); self.cur.children.append(n)
        if tag not in VOID: self.cur = n
    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag: n = n.parent
        if n is not self.root: self.cur = n.parent
    def handle_data(self, data):
        t = Node(None, parent=self.cur); t.text = data; self.cur.children.append(t)

html = open(SRC, encoding="utf-8").read()
html = re.sub(r"<style>.*?</style>", "", html, flags=re.S)
html = re.sub(r"<script>.*?</script>", "", html, flags=re.S)
b = Build(); b.feed(html); root = b.root

def gather(n):
    return (n.text or "") if n.tag is None else "".join(gather(c) for c in n.children)
def norm(s): return re.sub(r"[ \t\r\n]+", " ", s)

# ---------------- colours & doc ----------------
INK  = RGBColor(0x1A,0x1D,0x21); SOFT = RGBColor(0x5A,0x60,0x66)
FAINT= RGBColor(0x86,0x8C,0x92); ACC  = RGBColor(0x8C,0x2A,0x33)
ROLE_RGB = {"juge":RGBColor(0x2B,0x42,0x57),"proc":ACC,"def":RGBColor(0x1F,0x5D,0x52),
            "cun":RGBColor(0x6B,0x3A,0x63),"cand":RGBColor(0x7A,0x5A,0x1E)}
ROLE_TITLE = {"juge":"LE JUGE","proc":"LE PROCUREUR","def":"LA DÉFENSE","cun":"CUNÉGONDE","cand":"CANDIDE"}
ROLE_ELEVE = {"juge":"Élève 1","proc":"Élève 2","def":"Élève 3","cun":"Élève 4","cand":"Élève 5"}
KEYS = ["juge","proc","def","cun","cand"]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Garamond"; st.font.size = Pt(11.5)
st.font.color.rgb = INK; st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.16
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)

def rule(par, where="bottom", sz=4, color="D4D4CB"):
    pPr = par._p.get_or_add_pPr()
    bd = pPr.find(qn("w:pBdr"))
    if bd is None:
        bd = OxmlElement("w:pBdr"); pPr.append(bd)
    el = OxmlElement("w:"+where)
    el.set(qn("w:val"),"single"); el.set(qn("w:sz"),str(sz))
    el.set(qn("w:space"),"3"); el.set(qn("w:color"),color)
    bd.append(el)

def H(text, size, color=INK, before=16, after=5, sans=False, bold=False, caps=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper() if caps else text)
    r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold
    if sans: r.font.name = "Calibri"
    return p

# ---------------- inline rendering ----------------
def emit(par, node, bold=False, italic=False, color=None, size=None):
    for c in node.children:
        if c.tag is None:
            if not c.text: continue
            r = par.add_run(norm(c.text)); r.bold, r.italic = bold, italic
            r.font.color.rgb = color or INK
            if size: r.font.size = size
            continue
        cl = c.cls()
        if "beat" in cl: continue
        if c.tag == "br": par.add_run().add_break(); continue
        if "ref" in cl:
            r = par.add_run(" (" + norm(gather(c)).strip() + ")")
            r.font.size = Pt(8.5); r.font.color.rgb = SOFT; r.font.name = "Calibri"
            continue
        nb, ni, nc = bold, italic, color
        if c.tag in ("b","strong"): nb = True
        if c.tag in ("i","em"): ni = True
        if "cit" in cl: ni, nc = True, ACC
        if "di" in cl:  ni, nc = True, FAINT
        emit(par, c, nb, ni, nc, size)

# ---------------- collect the lines ----------------
def pause_for(di_text):
    t = di_text.lower()
    if "long silence" in t: return 3.5
    if "trois coups" in t:  return 3.0
    if "un temps" in t:     return 2.0
    if "mouvement" in t or "se rassoit" in t: return 2.0
    return 1.5

ACTS, LINES = [], []
for sec in root.find_all(lambda n: n.tag == "section" and "sec" in n.cls()):
    head = sec.first(lambda n: "sec-head" in n.cls())
    num = norm(gather(head.first(lambda n: "sec-num" in n.cls()))).strip() if head else ""
    title = norm(gather(head.first(lambda n: n.tag == "h2"))).strip() if head else ""
    secid = sec.attrs.get("id","")
    seclines = sec.find_all(lambda n: "line" in n.cls())
    if not seclines: continue
    act = {"num":num, "title":title, "id":secid, "lines":[]}
    for ln in seclines:
        role = ln.attrs.get("data-role","")
        note = ln.first(lambda n: "railnote" in n.cls())
        say  = ln.first(lambda n: "say" in n.cls())
        paras = [c for c in say.children if c.tag == "p"] if say else []
        flow  = [c for c in say.children if c.tag == "p" or "beat" in c.cls()] if say else []
        plain = norm(" ".join(gather(p) for p in paras)).strip()
        words = len(plain.split())
        dur = words / WPM * 60.0 + TRANSITION
        for di in (say.find_all(lambda n: "di" in n.cls()) if say else []):
            dur += pause_for(norm(gather(di)))
        rec = {"role":role, "note":norm(gather(note)).strip() if note else "",
               "paras":paras, "flow":flow, "plain":plain, "words":words, "dur":dur,
               "refs":len(ln.find_all(lambda n: "ref" in n.cls())), "act":act}
        act["lines"].append(rec); LINES.append(rec)
    ACTS.append(act)

t = 0.0
for i, L in enumerate(LINES, 1):
    L["n"] = i; L["start"] = t; t += L["dur"]; L["end"] = t
    if i == 1:
        L["cue_who"], L["cue_txt"] = "", "Vous ouvrez l'audience."
    else:
        prev = LINES[i-2]; w = prev["plain"].split()
        L["cue_who"] = ROLE_TITLE.get(prev["role"],"")
        L["cue_txt"] = ("… " if len(w) > 12 else "") + " ".join(w[-12:])
TOTAL = t

def mmss(sec): return "%02d:%02d" % (int(sec)//60, int(sec)%60)

# ---------------- cover ----------------
H("TRIBUNAL LITTÉRAIRE · VOLTAIRE, CANDIDE OU L'OPTIMISME (1759)", 8.5, SOFT, before=0, after=10, sans=True, caps=True)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
r = p.add_run("Le Procès de Cunégonde"); r.font.size = Pt(31); r.font.color.rgb = INK
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(12)
r = p.add_run("Script d'audience minuté — chaque réplique, dans l'ordre, avec son signal de départ.")
r.italic = True; r.font.size = Pt(12); r.font.color.rgb = SOFT
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14); rule(p)
r = p.add_run("Affaire n° 1759  ·  5 rôles  ·  %d répliques  ·  %d renvois au roman  ·  durée totale %s"
              % (len(LINES), sum(L["refs"] for L in LINES), mmss(TOTAL)))
r.font.size = Pt(9); r.font.color.rgb = SOFT; r.font.name = "Calibri"

H("COMMENT LIRE CE DOCUMENT", 9, ACC, before=10, after=4, sans=True, bold=True, caps=True)
for txt in [
 "Chaque prise de parole porte un numéro (N° 1 à %d) et un horaire de départ, compté depuis l'ouverture de l'audience. C'est le moment où vous devez parler." % len(LINES),
 "Sous le numéro, la ligne « Signal » donne les derniers mots de la personne qui parle juste avant vous. Dès que vous les entendez, c'est à vous. C'est ce qui remplace un metteur en scène.",
 "La note en italique à droite du nom est une consigne de jeu (ton, geste, silence) : elle ne se prononce pas. Les (indications entre parenthèses) à l'intérieur du texte ne se prononcent pas non plus, elles se jouent.",
 "Les passages en italique rouge sont des citations de Voltaire : dites-les mot pour mot. Les mentions (ch. I) indiquent le chapitre — c'est ce qui rapporte les 10 points de « références exactes au roman ».",
 "Le minutage est calculé à %d mots par minute, pauses et silences compris. Si vous parlez plus vite, retranchez environ 10 %%." % int(WPM),
]:
    q = doc.add_paragraph(); q.paragraph_format.left_indent = Inches(0.18); q.paragraph_format.space_after = Pt(4)
    r = q.add_run(txt); r.font.size = Pt(10.5)

# ---------------- running order ----------------
H("DÉROULÉ DE L'AUDIENCE", 9, ACC, before=18, after=5, sans=True, bold=True, caps=True)
rows = [["Acte","Titre","Début","Durée","Répliques"]]
for a in ACTS:
    s0 = a["lines"][0]["start"]; s1 = a["lines"][-1]["end"]
    rows.append([a["num"].replace("ACTE ","").strip(), a["title"], mmss(s0),
                 "%d min" % max(1, round((s1-s0)/60)),
                 "N° %d–%d" % (a["lines"][0]["n"], a["lines"][-1]["n"])])
tb = doc.add_table(rows=0, cols=5); tb.style = "Table Grid"
for i, row in enumerate(rows):
    cells = tb.add_row().cells
    for j, v in enumerate(row):
        cp = cells[j].paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
        rr = cp.add_run(v)
        rr.font.size = Pt(8.5) if i == 0 else Pt(10)
        if i == 0: rr.bold = True; rr.font.name = "Calibri"
for j, w in enumerate([0.55, 3.05, 0.75, 0.75, 1.05]):
    for row in tb.rows: row.cells[j].width = Inches(w)

doc.add_paragraph()
H("TEMPS DE PAROLE PAR ÉLÈVE", 9, ACC, before=14, after=5, sans=True, bold=True, caps=True)
rows = [["Élève","Rôle","Répliques","Références","Temps de parole"]]
for k in KEYS:
    mine = [L for L in LINES if L["role"] == k]
    rows.append([ROLE_ELEVE[k], ROLE_TITLE[k].title(), str(len(mine)),
                 str(sum(L["refs"] for L in mine)),
                 "%d min %02d s" % (int(sum(L["dur"] for L in mine))//60, int(sum(L["dur"] for L in mine))%60)])
tb = doc.add_table(rows=0, cols=5); tb.style = "Table Grid"
for i, row in enumerate(rows):
    cells = tb.add_row().cells
    for j, v in enumerate(row):
        cp = cells[j].paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
        rr = cp.add_run(v); rr.font.size = Pt(8.5) if i == 0 else Pt(10)
        if i == 0: rr.bold = True; rr.font.name = "Calibri"
for j, w in enumerate([0.8, 1.6, 0.95, 1.05, 1.45]):
    for row in tb.rows: row.cells[j].width = Inches(w)

# ---------------- part 1: full chronological script ----------------
def page_break():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

def render_line(L, show_time=True):
    hp = doc.add_paragraph()
    hp.paragraph_format.space_before = Pt(13); hp.paragraph_format.space_after = Pt(3)
    hp.paragraph_format.keep_with_next = True
    r = hp.add_run("N° %d" % L["n"]); r.bold = True; r.font.size = Pt(9); r.font.name = "Calibri"; r.font.color.rgb = FAINT
    if show_time:
        r = hp.add_run("   " + mmss(L["start"])); r.font.size = Pt(9); r.font.name = "Calibri"; r.font.color.rgb = FAINT
    r = hp.add_run("   " + ROLE_TITLE.get(L["role"], "")); r.bold = True; r.font.size = Pt(10.5)
    r.font.name = "Calibri"; r.font.color.rgb = ROLE_RGB.get(L["role"], INK)
    if L["note"]:
        r = hp.add_run("   " + L["note"]); r.italic = True; r.font.size = Pt(8.5)
        r.font.name = "Calibri"; r.font.color.rgb = FAINT
    rule(hp)
    cp = doc.add_paragraph(); cp.paragraph_format.left_indent = Inches(0.28)
    cp.paragraph_format.space_after = Pt(5); cp.paragraph_format.keep_with_next = True
    r = cp.add_run("Signal — "); r.font.size = Pt(8.5); r.font.name = "Calibri"; r.font.color.rgb = FAINT; r.bold = True
    if L["cue_who"]:
        r = cp.add_run(L["cue_who"] + " : "); r.font.size = Pt(8.5); r.font.name = "Calibri"; r.font.color.rgb = FAINT
    r = cp.add_run("« " + L["cue_txt"] + " »" if L["cue_who"] else L["cue_txt"])
    r.italic = True; r.font.size = Pt(8.5); r.font.name = "Calibri"; r.font.color.rgb = FAINT
    multi = len(L["paras"]) > 1
    t_in = L["start"]
    for node in L["flow"]:
        if node.tag != "p":                       # <span class="beat"> : change of movement
            sep = doc.add_paragraph()
            sep.paragraph_format.left_indent = Inches(0.78)
            sep.paragraph_format.space_before = Pt(4); sep.paragraph_format.space_after = Pt(7)
            rule(sep, color="C9C9BF", sz=6)
            continue
        par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(5)
        if multi:
            par.paragraph_format.left_indent = Inches(0.78)
            par.paragraph_format.first_line_indent = Inches(-0.50)
            par.paragraph_format.tab_stops.add_tab_stop(Inches(0.78), WD_TAB_ALIGNMENT.LEFT)
            r = par.add_run(mmss(t_in) + "\t")
            r.font.size = Pt(8); r.font.name = "Calibri"; r.font.color.rgb = FAINT
        else:
            par.paragraph_format.left_indent = Inches(0.78)
        emit(par, sp_ := node)
        w = len(norm(gather(node)).split())
        t_in += w / WPM * 60.0
        for di in node.find_all(lambda n: "di" in n.cls()):
            t_in += pause_for(norm(gather(di)))

page_break()
H("PREMIÈRE PARTIE", 9, ACC, before=0, after=2, sans=True, bold=True, caps=True)
H("Le procès intégral, dans l'ordre", 20, INK, before=0, after=4)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
r = p.add_run("Les %d répliques enchaînées. C'est le document que lit le juge et que suit la professeure." % len(LINES))
r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = SOFT

for a in ACTS:
    H("%s  ·  départ %s  ·  répliques N° %d à %d" % (a["num"], mmss(a["lines"][0]["start"]),
      a["lines"][0]["n"], a["lines"][-1]["n"]), 8.5, ACC, before=20, after=2, sans=True, caps=True)
    H(a["title"], 17, INK, before=0, after=3)
    for L in a["lines"]:
        render_line(L)

# ---------------- part 2: one sheet per student ----------------
page_break()
H("DEUXIÈME PARTIE", 9, ACC, before=0, after=2, sans=True, bold=True, caps=True)
H("Fiches individuelles", 20, INK, before=0, after=4)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
r = p.add_run("Une fiche par élève : uniquement vos répliques, dans l'ordre, avec l'heure de départ et le signal à écouter. "
              "Imprimez seulement votre fiche et apprenez-la ; vous n'avez pas besoin du texte des autres.")
r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = SOFT

for k in KEYS:
    mine = [L for L in LINES if L["role"] == k]
    page_break()
    H("FICHE — %s" % ROLE_ELEVE[k].upper(), 8.5, SOFT, before=0, after=2, sans=True, caps=True)
    H(ROLE_TITLE[k].title(), 24, ROLE_RGB[k], before=0, after=4)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(12); rule(p)
    secs = sum(L["dur"] for L in mine)
    r = p.add_run("%d répliques  ·  %d références au roman  ·  %d min %02d s de parole  ·  première entrée à %s"
                  % (len(mine), sum(L["refs"] for L in mine), int(secs)//60, int(secs)%60, mmss(mine[0]["start"])))
    r.font.size = Pt(9); r.font.color.rgb = SOFT; r.font.name = "Calibri"
    cur = None
    for L in mine:
        if L["act"] is not cur:
            cur = L["act"]
            H("%s — %s" % (cur["num"], cur["title"]), 9, ACC, before=14, after=2, sans=True, bold=True, caps=True)
        render_line(L)

# ---------------- part 3: annexes, carried over ----------------
page_break()
H("TROISIÈME PARTIE", 9, ACC, before=0, after=2, sans=True, bold=True, caps=True)
H("Annexes", 20, INK, before=0, after=10)

def add_html_table(tbl, widths):
    ths = [norm(gather(th)).strip() for th in tbl.find_all(lambda n: n.tag == "th")]
    t2 = doc.add_table(rows=0, cols=len(ths)); t2.style = "Table Grid"
    cells = t2.add_row().cells
    for j, v in enumerate(ths):
        cp = cells[j].paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
        rr = cp.add_run(v); rr.bold = True; rr.font.size = Pt(8.5); rr.font.name = "Calibri"
    for tr in tbl.find_all(lambda n: n.tag == "tr"):
        tds = tr.find_all(lambda n: n.tag == "td")
        if not tds: continue
        cells = t2.add_row().cells
        for j, td in enumerate(tds):
            cp = cells[j].paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
            emit(cp, td, size=Pt(9.5))
    for j, w in enumerate(widths):
        for row in t2.rows: row.cells[j].width = Inches(w)

for sec in root.find_all(lambda n: n.tag == "section" and "sec" in n.cls()):
    if not sec.attrs.get("id","").startswith(("annexes",)) and not sec.first(lambda n: "tip" in n.cls()) \
       and not (sec.first(lambda n: n.tag=="table") and not sec.find_all(lambda n: "line" in n.cls())):
        continue
    if sec.find_all(lambda n: "line" in n.cls()): continue
    head = sec.first(lambda n: "sec-head" in n.cls())
    if head:
        num = norm(gather(head.first(lambda n: "sec-num" in n.cls()))).strip()
        title = norm(gather(head.first(lambda n: n.tag == "h2"))).strip()
        tm = head.first(lambda n: "sec-time" in n.cls())
        H(num + ("  ·  " + norm(gather(tm)).strip() if tm else ""), 8.5, ACC, before=18, after=2, sans=True, caps=True)
        H(title, 17, INK, before=0, after=5)
    intro = [c for c in sec.children if c.tag == "p"]
    for ip in intro:
        par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(6)
        emit(par, ip, size=Pt(10.5))
        for rr in par.runs: rr.font.color.rgb = SOFT; rr.italic = True
    for tbl in sec.find_all(lambda n: n.tag == "table"):
        ncol = len(tbl.find_all(lambda n: n.tag == "th"))
        add_html_table(tbl, [0.8, 4.05, 1.65] if ncol == 3 else [2.2, 4.3])
        doc.add_paragraph()
    for tip in sec.find_all(lambda n: "tip" in n.cls()):
        h4 = tip.first(lambda n: n.tag == "h4")
        H(norm(gather(h4)).strip(), 10, INK, before=10, after=3, sans=True, bold=True, caps=True)
        for li in tip.find_all(lambda n: n.tag == "li"):
            par = doc.add_paragraph(style="List Bullet"); par.paragraph_format.space_after = Pt(2)
            emit(par, li, size=Pt(10))
    for nt in sec.find_all(lambda n: "note" in n.cls()):
        h4 = nt.first(lambda n: n.tag == "h4")
        title = norm(gather(h4)).strip() if h4 else ""
        if title.upper() == "MODE D'EMPLOI": continue
        if h4: H(title, 9, ACC, before=12, after=3, sans=True, bold=True, caps=True)
        for sp in [c for c in nt.children if c.tag == "p"]:
            par = doc.add_paragraph(); par.paragraph_format.left_indent = Inches(0.2)
            par.paragraph_format.space_after = Pt(4); emit(par, sp, size=Pt(10.5))

doc.save(OUT)
print("TOTAL %s  (%d répliques, %d mots, %d réfs)" % (mmss(TOTAL), len(LINES),
      sum(L["words"] for L in LINES), sum(L["refs"] for L in LINES)))
for k in KEYS:
    mine = [L for L in LINES if L["role"] == k]
    print("  %-6s %2d répliques  %2d réfs  %s de parole" % (k, len(mine), sum(L["refs"] for L in mine),
          mmss(sum(L["dur"] for L in mine))))
print("saved", OUT)
