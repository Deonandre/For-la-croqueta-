import io

import docx
import pytest
from fastapi.testclient import TestClient

from app import server


@pytest.fixture
def client(keyword_predictor, monkeypatch):
    server.get_predictor.cache_clear()
    monkeypatch.setattr(server, "get_predictor", lambda: keyword_predictor)
    return TestClient(server.app)


def test_index_and_health(client):
    assert client.get("/").status_code == 200
    assert client.get("/api/health").json()["ready"] is True


def test_analyze_text(client, essay_en):
    r = client.post("/api/analyze", json={"text": essay_en})
    assert r.status_code == 200
    body = r.json()
    assert body["text"] == essay_en
    assert body["language"] == "en"
    assert body["segments"] and body["summary"]["verdict"] in {"human", "mostly_human"}


def test_analyze_docx(client):
    d = docx.Document()
    d.add_paragraph("Introduction")
    d.add_paragraph("It is crucial to note this crucial point. " * 5)
    table = d.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Year"
    buf = io.BytesIO()
    d.save(buf)
    r = client.post("/api/analyze-file", files={"file": ("ee.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200
    assert "Introduction" in r.json()["text"]
    assert r.json()["passages"][0]["label"] == "ai"


def test_errors(client):
    assert client.post("/api/analyze", json={"text": "   "}).json()["detail"] == "empty_text"
    r = client.post("/api/analyze-file", files={"file": ("x.exe", b"123", "application/octet-stream")})
    assert r.status_code == 415


def test_no_model_configured(monkeypatch):
    monkeypatch.delenv("AIDETECT_MODEL", raising=False)
    server.get_predictor.cache_clear()
    c = TestClient(server.app)
    assert c.get("/api/health").json() == {"ready": False}
    assert c.post("/api/analyze", json={"text": "hello world"}).json()["detail"] == "no_model"
